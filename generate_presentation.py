import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    PRIMARY_COLOR = RGBColor(180, 50, 20)       # Warm Terracotta
    PRIMARY_DARK = RGBColor(120, 25, 10)       # Dark Burgundy
    SECONDARY_COLOR = RGBColor(220, 100, 40)   # Tangerine Amber
    ACCENT_COLOR = RGBColor(245, 166, 35)      # Golden Turmeric
    BG_LIGHT = RGBColor(250, 248, 245)         # Warm Off-white
    CARD_BG = RGBColor(255, 255, 255)          # White
    TEXT_MAIN = RGBColor(35, 31, 32)           # Charcoal
    TEXT_MUTED = RGBColor(100, 100, 105)       # Grey
    BORDER_COLOR = RGBColor(228, 222, 215)
    SUCCESS_COLOR = RGBColor(34, 139, 34)
    INFO_BLUE = RGBColor(25, 118, 210)

    blank_layout = prs.slide_layouts[6]

    def set_slide_background(slide, color):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
        bg.fill.solid()
        bg.fill.fore_color.rgb = color
        bg.line.fill.background()
        return bg

    def add_header(slide, title_text, category_text="WEBSITE RESEP MASAKAN • PROYEK WEB"):
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        tf_cat.margin_left = tf_cat.margin_top = tf_cat.margin_right = tf_cat.margin_bottom = 0
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = SECONDARY_COLOR

        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.75), Inches(11.7), Inches(0.6))
        tf = title_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.size = Pt(24)
        p.font.bold = True
        p.font.color.rgb = PRIMARY_DARK

    def add_card(slide, left, top, width, height, bg_color=CARD_BG, border_color=BORDER_COLOR):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        if border_color:
            card.line.color.rgb = border_color
            card.line.width = Pt(1.2)
        else:
            card.line.fill.background()
        return card

    # ==========================================
    # SLIDE 1: Cover + List Nama Anggota Kecil
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide1, PRIMARY_DARK)

    t_box = slide1.shapes.add_textbox(Inches(1.2), Inches(1.1), Inches(11.0), Inches(2.6))
    tf1 = t_box.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "PRESENTASI PROYEK WEBSITE"
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_COLOR
    p0.space_after = Pt(10)

    p1 = tf1.add_paragraph()
    p1.text = "🍳 WEBSITE RESEP MASAKAN"
    p1.font.size = Pt(40)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(255, 255, 255)
    p1.space_after = Pt(10)

    p2 = tf1.add_paragraph()
    p2.text = "Platform Resep Masakan Nusantara Modern berbasis DATH Stack\n(Django 5.2 • Alpine.js 3 • Tailwind CSS • HTMX 2)"
    p2.font.size = Pt(18)
    p2.font.color.rgb = RGBColor(245, 235, 225)

    # Box Anggota Kelompok (List Ukuran Lebih Kecil)
    group_card = add_card(slide1, Inches(1.2), Inches(4.2), Inches(6.5), Inches(2.4), bg_color=RGBColor(70, 18, 8), border_color=RGBColor(180, 70, 40))
    gb_box = slide1.shapes.add_textbox(Inches(1.4), Inches(4.35), Inches(6.1), Inches(2.1))
    tf_g = gb_box.text_frame
    tf_g.word_wrap = True

    pg0 = tf_g.paragraphs[0]
    pg0.text = "ANGGOTA KELOMPOK:"
    pg0.font.size = Pt(11)
    pg0.font.bold = True
    pg0.font.color.rgb = ACCENT_COLOR
    pg0.space_after = Pt(6)

    members = ["•  Ali Alhafidz", "•  Aldiana Risman", "•  Riski Hidayat", "•  Miqdad Alfarizi"]
    for m in members:
        pg = tf_g.add_paragraph()
        pg.text = m
        pg.font.size = Pt(12)
        pg.font.color.rgb = RGBColor(255, 245, 235)
        pg.space_after = Pt(3)

    # ==========================================
    # SLIDE 2: Latar Belakang & Solusi (Satu Paragraf Masing-Masing)
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide2, BG_LIGHT)
    add_header(slide2, "Latar Belakang & Solusi", "02 • LATAR BELAKANG PROYEK")

    # Card 1: Latar Belakang (Satu Paragraf)
    card1_left = Inches(0.8)
    card_width = Inches(5.65)
    card_top = Inches(1.6)
    card_height = Inches(5.1)

    add_card(slide2, card1_left, card_top, card_width, card_height)
    bar1 = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, card1_left, card_top, card_width, Inches(0.12))
    bar1.fill.solid()
    bar1.fill.fore_color.rgb = PRIMARY_COLOR
    bar1.line.fill.background()

    tb1 = slide2.shapes.add_textbox(card1_left + Inches(0.4), card_top + Inches(0.4), card_width - Inches(0.8), card_height - Inches(0.8))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p_bg_title = tf1.paragraphs[0]
    p_bg_title.text = "⚠️ Latar Belakang Masalah"
    p_bg_title.font.size = Pt(20)
    p_bg_title.font.bold = True
    p_bg_title.font.color.rgb = PRIMARY_COLOR
    p_bg_title.space_after = Pt(18)

    p_bg_desc = tf1.add_paragraph()
    p_bg_desc.text = "Kekayaan kuliner tradisional nusantara saat ini masih banyak tersebar secara tidak terstruktur dengan takaran yang seringkali tidak baku, sehingga menyulitkan koki rumahan dalam memperkirakan takaran bahan saat porsi diubah dan mencatat daftar belanja secara praktis, diperparah oleh kebanyakan website resep modern yang lambat serta berat diakses karena beban framework JavaScript yang berlebihan."
    p_bg_desc.font.size = Pt(15)
    p_bg_desc.font.color.rgb = TEXT_MAIN
    p_bg_desc.space_after = Pt(14)

    # Card 2: Solusi (Satu Paragraf)
    card2_left = Inches(6.88)
    add_card(slide2, card2_left, card_top, card_width, card_height)
    bar2 = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, card2_left, card_top, card_width, Inches(0.12))
    bar2.fill.solid()
    bar2.fill.fore_color.rgb = SUCCESS_COLOR
    bar2.line.fill.background()

    tb2 = slide2.shapes.add_textbox(card2_left + Inches(0.4), card_top + Inches(0.4), card_width - Inches(0.8), card_height - Inches(0.8))
    tf2 = tb2.text_frame
    tf2.word_wrap = True

    p_sol_title = tf2.paragraphs[0]
    p_sol_title.text = "💡 Solusi Web Masakan Ini"
    p_sol_title.font.size = Pt(20)
    p_sol_title.font.bold = True
    p_sol_title.font.color.rgb = SUCCESS_COLOR
    p_sol_title.space_after = Pt(18)

    p_sol_desc = tf2.add_paragraph()
    p_sol_desc.text = "Website resep masakan ini menghadirkan platform katalog resep nusantara terstandarisasi berbasis arsitektur DATH Stack yang sangat ringan dan cepat, dilengkapi fitur cerdas kalkulator porsi dinamis untuk penyesuaian takaran bahan secara otomatis serta integrasi daftar belanja digital yang mempermudah seluruh proses memasak hidangan tradisional dari awal hingga akhir."
    p_sol_desc.font.size = Pt(15)
    p_sol_desc.font.color.rgb = TEXT_MAIN
    p_sol_desc.space_after = Pt(14)

    # ==========================================
    # SLIDE 3: DATH Stack
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide3, BG_LIGHT)
    add_header(slide3, "Arsitektur Teknologi: DATH Stack", "03 • STACK TEKNOLOGI")

    tech_stack = [
        ("Django 5.2", "Core Backend & Security", "ORM handal, manajemen autentikasi & peran pengguna, serta proteksi CSRF & SQL Injection."),
        ("HTMX 2.x", "Interaktivitas Tanpa Reload", "Pencarian instan, filter kategori, dan tombol simpan favorit tanpa refresh halaman penuh."),
        ("Alpine.js 3.x", "Reaktivitas Mikro Klien", "Slider kalkulasi porsi piring real-time, drawer mobile, dan dialog modal langsung di browser."),
        ("Tailwind CSS", "Material Design 3", "Desain berbasis utility-first dengan palet warna hangat khas rempah yang responsif dan mobile-friendly.")
    ]
    for i, (t_name, t_role, t_desc) in enumerate(tech_stack):
        c_left = Inches(0.8 + (i % 2) * 5.95)
        c_top = Inches(1.6 + (i // 2) * 2.65)
        c_width = Inches(5.75)
        c_height = Inches(2.4)
        add_card(slide3, c_left, c_top, c_width, c_height)

        tb = slide3.shapes.add_textbox(c_left + Inches(0.3), c_top + Inches(0.25), c_width - Inches(0.6), c_height - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = t_name
        p1.font.size = Pt(18)
        p1.font.bold = True
        p1.font.color.rgb = PRIMARY_COLOR

        p2 = tf.add_paragraph()
        p2.text = t_role
        p2.font.size = Pt(12)
        p2.font.bold = True
        p2.font.color.rgb = SECONDARY_COLOR
        p2.space_after = Pt(8)

        p3 = tf.add_paragraph()
        p3.text = t_desc
        p3.font.size = Pt(12)
        p3.font.color.rgb = TEXT_MAIN

    # ==========================================
    # SLIDE 4: Fitur Pengguna
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide4, BG_LIGHT)
    add_header(slide4, "Fitur Unggulan Pengguna", "04 • FITUR APLIKASI")

    features_user = [
        ("🔍 Cari & Filter Cepat", "Filter resep berdasarkan bahan utama, daerah, waktu masak, dan tingkat kesulitan secara instan."),
        ("⚖️ Porsi Dinamis", "Kalkulasi otomatis jumlah takaran bahan (gram, sdm) saat jumlah porsi piring diubah pengguna."),
        ("🛒 Daftar Belanja Cerdas", "Tambahkan bahan ke shopping list pribadi dengan penggabungan item sejenis secara otomatis."),
        ("❤️ Koleksi Favorit", "Simpan resep pilihan pengguna dengan satu klik tanpa reload halaman."),
        ("🥗 Estimasi Gizi", "Rincian kalori, protein, lemak, dan karbohidrat per porsi untuk panduan hidup sehat."),
        ("✍️ Berbagi Resep", "Form pembuatan resep lengkap dengan formset dinamis langkah masak dan foto.")
    ]
    for i, (f_title, f_desc) in enumerate(features_user):
        c_left = Inches(0.8 + (i % 3) * 3.98)
        c_top = Inches(1.6 + (i // 3) * 2.65)
        c_width = Inches(3.78)
        c_height = Inches(2.4)
        add_card(slide4, c_left, c_top, c_width, c_height)

        tb = slide4.shapes.add_textbox(c_left + Inches(0.25), c_top + Inches(0.25), c_width - Inches(0.5), c_height - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = f_title
        p1.font.size = Pt(15)
        p1.font.bold = True
        p1.font.color.rgb = PRIMARY_DARK
        p1.space_after = Pt(8)
        p2 = tf.add_paragraph()
        p2.text = f_desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = TEXT_MAIN

    # ==========================================
    # SLIDE 5: Moderasi & Admin
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide5, BG_LIGHT)
    add_header(slide5, "Moderasi Resep & Dashboard Admin", "05 • TATA KELOLA KONTEN")

    admin_cols = [
        ("Alur Moderasi Resep",
         "• Resep baru berstatus 'PENDING' (belum publik).\n• Admin memverifikasi dan menyetujui (PUBLISHED).\n• Opsi penolakan wajib menyertakan alasan yang jelas.\n• Penulis tetap dapat membaca catatan perbaikan.",
         PRIMARY_COLOR),
        ("Dashboard Administrator",
         "• Metrik total resep terbit, antrean pending, dan user.\n• Aksi cepat moderasi langsung dari tabel review.\n• Kontrol taksonomi kategori dan tag pencarian.\n• Proteksi aksi menggunakan metode POST dan CSRF.",
         SECONDARY_COLOR),
        ("Pemisahan Hak Akses (Role)",
         "• Role Pengguna Biasa: Membaca, menulis resep, belanja, simpan favorit.\n• Role Admin: Manajemen konten, kurasi, dan moderasi komunitas.",
         SUCCESS_COLOR)
    ]
    card_w = Inches(3.68)
    card_h = Inches(5.1)
    card_top = Inches(1.6)

    for i, (title, content, bar_color) in enumerate(admin_cols):
        cleft = Inches(0.8 + i * 4.02)
        add_card(slide5, cleft, card_top, card_w, card_h)
        bar = slide5.shapes.add_shape(MSO_SHAPE.RECTANGLE, cleft, card_top, card_w, Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = bar_color
        bar.line.fill.background()

        tb = slide5.shapes.add_textbox(cleft + Inches(0.25), card_top + Inches(0.3), card_w - Inches(0.5), card_h - Inches(0.5))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(17)
        p.font.bold = True
        p.font.color.rgb = bar_color
        p.space_after = Pt(14)
        for line in content.split("\n"):
            p_line = tf.add_paragraph()
            p_line.text = line
            p_line.font.size = Pt(13)
            p_line.font.color.rgb = TEXT_MAIN
            p_line.space_after = Pt(10)

    # ==========================================
    # SLIDE 6: Struktur Database
    # ==========================================
    slide6 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide6, BG_LIGHT)
    add_header(slide6, "Struktur Database Relasional", "06 • DATABASE ARCHITECTURE")

    models_data = [
        ("User & Profile", "Pemisahan peran Admin dan Pengguna, profil bio koki dan avatar."),
        ("Category & Tag", "Klasifikasi daerah nusantara (Padang, Sunda) dan karakteristik masakan."),
        ("Recipe (Utama)", "Menyimpan judul, slug unik, estimasi waktu masak, porsi, dan status publikasi."),
        ("Ingredient & Step", "Bahan masakan dengan takaran angka Decimal serta urutan langkah instruksi."),
        ("Shopping List", "Daftar belanja bahan milik pengguna dengan fitur penggabungan otomatis."),
        ("Favorite", "Menyimpan resep-resep pilihan yang disukai oleh pengguna.")
    ]
    for i, (m_title, m_desc) in enumerate(models_data):
        c_left = Inches(0.8 + (i % 3) * 3.98)
        c_top = Inches(1.6 + (i // 3) * 2.65)
        c_width = Inches(3.78)
        c_height = Inches(2.4)
        add_card(slide6, c_left, c_top, c_width, c_height)

        tb = slide6.shapes.add_textbox(c_left + Inches(0.25), c_top + Inches(0.25), c_width - Inches(0.5), c_height - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        p1 = tf.paragraphs[0]
        p1.text = m_title
        p1.font.size = Pt(16)
        p1.font.bold = True
        p1.font.color.rgb = PRIMARY_DARK
        p1.space_after = Pt(8)
        p2 = tf.add_paragraph()
        p2.text = m_desc
        p2.font.size = Pt(12)
        p2.font.color.rgb = TEXT_MAIN

    # ==========================================
    # SLIDE 7: Roadmap Masa Depan (Slide 7)
    # ==========================================
    slide7 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide7, BG_LIGHT)
    add_header(slide7, "Rencana Pengembangan Masa Depan (Roadmap)", "07 • ROADMAP & INOVASI")

    roadmap_items = [
        ("Tahap 1: Komunitas", "• Sistem rating bintang & ulasan resep.\n• Unggah foto hasil memasak (Cooksnap).\n• Buku koleksi resep kustom pengguna.", PRIMARY_COLOR),
        ("Tahap 2: AI Smart Kitchen", "• Rekomendasi menu dari bahan kulkas.\n• Ekstraksi resep dari foto buku masak.\n• Timer memasak interaktif per langkah.", SECONDARY_COLOR),
        ("Tahap 3: Aksesibilitas", "• Mode offline PWA tanpa kuota internet.\n• Ekspor shopping list ke WhatsApp.\n• Dukungan multi-bahasa.", SUCCESS_COLOR)
    ]
    card_h_s7 = Inches(4.0)
    for i, (r_title, r_desc, accent) in enumerate(roadmap_items):
        cleft = Inches(0.8 + i * 4.02)
        add_card(slide7, cleft, card_top, card_w, card_h_s7)
        bar = slide7.shapes.add_shape(MSO_SHAPE.RECTANGLE, cleft, card_top, card_w, Inches(0.12))
        bar.fill.solid()
        bar.fill.fore_color.rgb = accent
        bar.line.fill.background()

        tb = slide7.shapes.add_textbox(cleft + Inches(0.25), card_top + Inches(0.25), card_w - Inches(0.5), card_h_s7 - Inches(0.4))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = r_title
        p.font.size = Pt(17)
        p.font.bold = True
        p.font.color.rgb = accent
        p.space_after = Pt(12)
        for line in r_desc.split("\n"):
            p_line = tf.add_paragraph()
            p_line.text = line
            p_line.font.size = Pt(12)
            p_line.font.color.rgb = TEXT_MAIN
            p_line.space_after = Pt(8)

    # Tombol Demo Web Lokal di Slide 7
    btn_box = slide7.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(3.8), Inches(5.95), Inches(5.7), Inches(0.85))
    btn_box.fill.solid()
    btn_box.fill.fore_color.rgb = PRIMARY_COLOR
    btn_box.line.fill.background()

    tf_btn = btn_box.text_frame
    tf_btn.word_wrap = True
    p_btn = tf_btn.paragraphs[0]
    p_btn.alignment = PP_ALIGN.CENTER
    run_btn = p_btn.add_run()
    run_btn.text = "🌐 BUKA DEMO WEB LOKAL (127.0.0.1:8000)"
    run_btn.font.size = Pt(14)
    run_btn.font.bold = True
    run_btn.font.color.rgb = RGBColor(255, 255, 255)
    run_btn.hyperlink.address = "http://127.0.0.1:8000/"

    # ==========================================
    # SLIDE 8: Kesimpulan (Satu Paragraf Bersih)
    # ==========================================
    slide8 = prs.slides.add_slide(blank_layout)
    set_slide_background(slide8, PRIMARY_DARK)

    t_box8 = slide8.shapes.add_textbox(Inches(1.5), Inches(1.6), Inches(10.33), Inches(4.5))
    tf8 = t_box8.text_frame
    tf8.word_wrap = True

    p0 = tf8.paragraphs[0]
    p0.text = "08 • KESIMPULAN PROYEK"
    p0.font.size = Pt(14)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_COLOR
    p0.space_after = Pt(16)

    p1 = tf8.add_paragraph()
    p1.text = "Kesimpulan"
    p1.font.size = Pt(36)
    p1.font.bold = True
    p1.font.color.rgb = RGBColor(255, 255, 255)
    p1.space_after = Pt(24)

    # Satu Paragraf
    p2 = tf8.add_paragraph()
    p2.text = "Website resep masakan ini berhasil menghadirkan platform kuliner nusantara yang ringan, cepat, dan interaktif berkat perpaduan arsitektur DATH Stack (Django, Alpine.js, Tailwind CSS, HTMX), dilengkapi fitur cerdas seperti kalkulator porsi dinamis dan daftar belanja otomatis yang memberikan solusi praktis, efisien, serta pengalaman memasak yang menyenangkan bagi seluruh keluarga."
    p2.font.size = Pt(17)
    p2.font.color.rgb = RGBColor(245, 235, 225)
    p2.space_after = Pt(32)

    p3 = tf8.add_paragraph()
    p3.text = "Terima Kasih! • Ada Pertanyaan?"
    p3.font.size = Pt(22)
    p3.font.bold = True
    p3.font.color.rgb = ACCENT_COLOR

    output_path = r"c:\Users\LAB1_CLIENT25\Desktop\resep_makan\Presentasi_Resep_Masakan.pptx"
    prs.save(output_path)
    print(f"Presentation saved successfully to {output_path}")

if __name__ == "__main__":
    create_presentation()

